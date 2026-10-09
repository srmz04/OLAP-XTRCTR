"""
ResultsScreen: muestra resultados de consulta MDX en tabla eficiente.
Modelo propio para manejar miles de filas sin copiar a QTableWidget.
Exporta a CSV y XLSX (si openpyxl disponible).
"""
from __future__ import annotations

import csv
import os
import time
from typing import Any, Dict, List, Optional

from PyQt5.QtCore import (
    QAbstractTableModel,
    QModelIndex,
    QSortFilterProxyModel,
    Qt,
    pyqtSignal,
    pyqtSlot,
)
from PyQt5.QtGui import QColor, QFont
from PyQt5.QtWidgets import (
    QApplication,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QPushButton,
    QSizePolicy,
    QTableView,
    QVBoxLayout,
    QWidget,
)

from core.provider import QueryResult
from ui import theme


# ── Custom table model ────────────────────────────────────────────────────────

class _OlapTableModel(QAbstractTableModel):
    """Modelo eficiente: no copia datos, referencia directa a list-of-dicts."""

    def __init__(self, result: QueryResult, parent=None):
        super().__init__(parent)
        self._rows = result.rows
        self._cols = [c.get("headerName", c.get("field", f"Col{i}"))
                      for i, c in enumerate(result.columns)]
        self._fields = [c.get("field", c.get("headerName", f"Col{i}"))
                        for i, c in enumerate(result.columns)]

    def rowCount(self, parent=QModelIndex()) -> int:
        return len(self._rows)

    def columnCount(self, parent=QModelIndex()) -> int:
        return len(self._cols)

    def data(self, index: QModelIndex, role=Qt.DisplayRole):
        if not index.isValid():
            return None
        r, c = index.row(), index.column()
        if r >= len(self._rows) or c >= len(self._fields):
            return None

        val = self._rows[r].get(self._fields[c], "")

        if role == Qt.DisplayRole:
            if val is None:
                return ""
            return str(val)

        if role == Qt.TextAlignmentRole:
            # Numeric values right-aligned
            if isinstance(val, (int, float)):
                return Qt.AlignRight | Qt.AlignVCenter
            return Qt.AlignLeft | Qt.AlignVCenter

        if role == Qt.ForegroundRole:
            if isinstance(val, (int, float)):
                return QColor(theme.ACCENT_FG)
            return QColor(theme.TEXT)

        if role == Qt.BackgroundRole:
            if r % 2 == 1:
                return QColor(theme.BG_ALT)
            return QColor(theme.BG_BASE)

        return None

    def headerData(self, section: int, orientation: Qt.Orientation,
                   role=Qt.DisplayRole):
        if role == Qt.DisplayRole:
            if orientation == Qt.Horizontal:
                if section < len(self._cols):
                    return self._cols[section]
            else:
                return str(section + 1)
        if role == Qt.ForegroundRole:
            return QColor(theme.TEXT_DIM)
        return None

    def get_all_data(self) -> tuple:
        return self._cols, self._rows, self._fields


# ── Screen ─────────────────────────────────────────────────────────────────────

class ResultsScreen(QWidget):
    """
    Señales:
      new_query()   — volver al builder
      go_explorer() — volver al explorador
    """
    new_query = pyqtSignal()
    go_explorer = pyqtSignal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self._result: Optional[QueryResult] = None
        self._model: Optional[_OlapTableModel] = None
        self._elapsed: float = 0.0
        self._catalog: str = ""
        self._mdx: str = ""
        self._build_ui()

    def show_result(self, result: QueryResult, catalog: str = "",
                    mdx: str = "", elapsed: float = 0.0):
        self._result = result
        self._catalog = catalog
        self._mdx = mdx
        self._elapsed = elapsed
        self._populate(result)

    # ── UI ────────────────────────────────────────────────────────────────────

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        # Top bar
        topbar = QFrame()
        topbar.setFixedHeight(52)
        topbar.setProperty("class", "topbar")
        tb = QHBoxLayout(topbar)
        tb.setContentsMargins(16, 0, 16, 0)

        back_btn = QPushButton("← Nueva consulta")
        back_btn.clicked.connect(self.new_query)
        back_btn.setProperty("class", "ghost")
        tb.addWidget(back_btn)

        tb.addStretch()

        title = QLabel("Resultados")
        title.setStyleSheet(f"color: {theme.TEXT}; font-size: 15px; font-weight: 700;")
        tb.addWidget(title)

        tb.addStretch()

        explorer_btn = QPushButton("Explorador")
        explorer_btn.clicked.connect(self.go_explorer)
        explorer_btn.setStyleSheet(
            f"QPushButton {{ background: transparent; border: 1px solid {theme.BORDER};"
            f"color: {theme.TEXT_DIM}; border-radius: 6px; padding: 4px 14px; font-size: 12px; }}"
            f"QPushButton:hover {{ border-color: {theme.ACCENT}; color: {theme.ACCENT_FG}; }}"
        )
        tb.addWidget(explorer_btn)

        root.addWidget(topbar)

        # Info bar
        self._info_bar = QFrame()
        self._info_bar.setFixedHeight(40)
        self._info_bar.setStyleSheet(f"background: {theme.BG_DEEP}; border-bottom: 1px solid {theme.BORDER};")
        ib = QHBoxLayout(self._info_bar)
        ib.setContentsMargins(16, 0, 16, 0)
        ib.setSpacing(20)

        self._rows_lbl = QLabel("0 filas")
        self._rows_lbl.setStyleSheet(f"color: {theme.SUCCESS}; font-weight: 600; font-size: 12px;")
        ib.addWidget(self._rows_lbl)

        self._cols_lbl = QLabel("0 columnas")
        self._cols_lbl.setStyleSheet(f"color: {theme.GROUP_DIM}; font-size: 12px;")
        ib.addWidget(self._cols_lbl)

        self._time_lbl = QLabel("")
        self._time_lbl.setStyleSheet(f"color: {theme.TEXT_DIM}; font-size: 12px;")
        ib.addWidget(self._time_lbl)

        self._trunc_lbl = QLabel("")
        self._trunc_lbl.setStyleSheet(f"color: {theme.WARN}; font-size: 12px;")
        ib.addWidget(self._trunc_lbl)

        ib.addStretch()

        csv_btn = QPushButton("↓ CSV")
        csv_btn.setFixedHeight(28)
        csv_btn.setProperty("class", "success")
        csv_btn.clicked.connect(self._export_csv)
        ib.addWidget(csv_btn)

        xlsx_btn = QPushButton("↓ XLSX")
        xlsx_btn.setFixedHeight(28)
        xlsx_btn.setProperty("class", "success")
        xlsx_btn.clicked.connect(self._export_xlsx)
        ib.addWidget(xlsx_btn)

        root.addWidget(self._info_bar)

        # Table
        self._table = QTableView()
        self._table.setSortingEnabled(True)
        self._table.setAlternatingRowColors(True)
        self._table.horizontalHeader().setStretchLastSection(True)
        self._table.horizontalHeader().setMinimumSectionSize(80)
        self._table.verticalHeader().setDefaultSectionSize(28)
        self._table.setSelectionBehavior(QTableView.SelectRows)
        self._table.setFrameShape(QFrame.NoFrame)
        root.addWidget(self._table, 1)

        # Empty state overlay
        self._empty_lbl = QLabel("Sin resultados.\nEjecuta una consulta para ver datos.")
        self._empty_lbl.setAlignment(Qt.AlignCenter)
        self._empty_lbl.setStyleSheet(
            f"color: {theme.TEXT_MUTED}; font-size: 16px; background: transparent;"
        )
        self._empty_lbl.setVisible(False)
        root.addWidget(self._empty_lbl)

    # ── Populate ──────────────────────────────────────────────────────────────

    def _populate(self, result: QueryResult):
        if not result.rows:
            self._show_empty(result)
            return

        self._empty_lbl.setVisible(False)
        self._table.setVisible(True)

        model = _OlapTableModel(result)
        proxy = QSortFilterProxyModel()
        proxy.setSourceModel(model)
        self._table.setModel(proxy)
        self._model = model

        # Resize columns to content (up to 250px)
        for i in range(model.columnCount()):
            self._table.resizeColumnToContents(i)
            if self._table.columnWidth(i) > 250:
                self._table.setColumnWidth(i, 250)

        # Update info bar
        n_rows = result.row_count
        n_cols = model.columnCount()
        self._rows_lbl.setText(f"{n_rows:,} filas")
        self._cols_lbl.setText(f"{n_cols} columnas")
        if self._elapsed:
            self._time_lbl.setText(f"{self._elapsed:.2f}s")
        if result.truncated:
            self._trunc_lbl.setText("⚠ Resultado truncado")
        else:
            self._trunc_lbl.setText("")

    def _show_empty(self, result: QueryResult):
        self._table.setModel(None)
        self._table.setVisible(False)
        self._empty_lbl.setVisible(True)
        self._rows_lbl.setText("0 filas")
        self._cols_lbl.setText(f"{len(result.columns)} columnas")
        if self._elapsed:
            self._time_lbl.setText(f"{self._elapsed:.2f}s")
        self._trunc_lbl.setText("")

    # ── Exports ───────────────────────────────────────────────────────────────

    @pyqtSlot()
    def _export_csv(self):
        if not self._model:
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Exportar CSV", "resultados.csv",
            "CSV (*.csv);;Todos los archivos (*)"
        )
        if not path:
            return
        try:
            cols, rows, fields = self._model.get_all_data()
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=fields)
                writer.writeheader()
                for row in rows:
                    writer.writerow({k: row.get(k, "") for k in fields})
            QMessageBox.information(self, "Exportado", f"CSV guardado en:\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "Error exportando", str(exc))

    @pyqtSlot()
    def _export_xlsx(self):
        if not self._model:
            return
        try:
            import openpyxl  # type: ignore
        except ImportError:
            QMessageBox.warning(
                self, "openpyxl no disponible",
                "Instala openpyxl para exportar XLSX:\n  pip install openpyxl"
            )
            return

        path, _ = QFileDialog.getSaveFileName(
            self, "Exportar XLSX", "resultados.xlsx",
            "Excel (*.xlsx);;Todos los archivos (*)"
        )
        if not path:
            return
        try:
            cols, rows, fields = self._model.get_all_data()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Resultados"
            ws.append(cols)
            for row in rows:
                ws.append([row.get(f, "") for f in fields])
            # Header style
            from openpyxl.styles import Font, PatternFill, Alignment
            hdr_font = Font(bold=True, color=theme.TEXT.lstrip("#"))
            hdr_fill = PatternFill("solid", fgColor=theme.ACCENT.lstrip("#"))
            for cell in ws[1]:
                cell.font = hdr_font
                cell.fill = hdr_fill
                cell.alignment = Alignment(horizontal="center")
            wb.save(path)
            QMessageBox.information(self, "Exportado", f"XLSX guardado en:\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "Error exportando", str(exc))
